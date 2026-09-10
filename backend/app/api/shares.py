
# In-memory debounce cache to prevent duplicate view increments (IP + ShareID within 5 seconds)
_VIEW_DEBOUNCE_CACHE = {}

"""
SecureCloud - Shared Links & Public Access API
Expiration countdowns, password protection, global revocation, and format-preserving public views.
"""

from datetime import datetime, timedelta
from typing import Dict, Any, List, Optional
from fastapi import APIRouter, Depends, HTTPException, status, Request, Body
from fastapi.responses import FileResponse as FastAPIFileResponse, StreamingResponse
from sqlalchemy.orm import Session
import bcrypt
import io
import os

from backend.app.database import get_db
from backend.app.models.models import User, FileRecord, SharedLink
from backend.app.schemas.schemas import CreateSharedLinkRequest, SharedLinkResponse, VerifySharePasswordRequest
from backend.app.security.auth_utils import get_current_user, get_current_admin
from backend.app.services.audit_service import AuditService
from backend.app.services.storage_service import format_size, resolve_storage_path, sanitize_filename
from backend.app.services.event_service import EventService

def to_ist(dt: Optional[datetime]) -> str:
    if not dt:
        return "Never"
    ist_time = dt + timedelta(hours=5, minutes=30)
    return ist_time.strftime("%d %b %Y, %I:%M:%S %p IST")

def to_ist_short(dt: Optional[datetime]) -> str:
    if not dt:
        return "Never"
    ist_time = dt + timedelta(hours=5, minutes=30)
    return ist_time.strftime("%d %b %Y, %I:%M %p IST")

router = APIRouter(prefix="/api/shares", tags=["Shared Links"])

def _process_create_share(file_id: str, req: CreateSharedLinkRequest, current_user: User, db: Session):
    f = db.query(FileRecord).filter(
        FileRecord.id == file_id,
        FileRecord.user_id == current_user.id,
        FileRecord.is_in_recycle_bin == False
    ).first()

    if not f:
        raise HTTPException(status_code=404, detail="File not found or cannot be shared.")

    if f.is_confidential:
        raise HTTPException(status_code=400, detail="Confidential vault files cannot be publicly shared.")

    # Calculate expiration
    expires_at = None
    if req.expires_in_hours:
        expires_at = datetime.utcnow() + timedelta(hours=req.expires_in_hours)

    # Password hash
    pw_hash = None
    if req.password:
        pw_hash = bcrypt.hashpw(req.password.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')

    share = SharedLink(
        user_id=current_user.id,
        file_id=f.id,
        password_hash=pw_hash,
        is_password_protected=bool(req.password),
        expires_at=expires_at,
        view_only=req.view_only,
        allow_download=req.allow_download,
        is_active=True,
        view_count=0,
        download_count=0,
        created_at=datetime.utcnow()
    )
    db.add(share)
    db.commit()
    db.refresh(share)

    AuditService.log(
        db, "CREATE_SHARED_LINK", f"File '{f.filename}'", "SUCCESS",
        f"Created share link. Expiry: {req.expires_in_hours} hrs, Password: {bool(req.password)}",
        user_id=current_user.id, username=current_user.username, role=current_user.role
    )

    return {
        "id": share.id,
        "file_id": f.id,
        "filename": f.filename,
        "file_size": f.file_size,
        "file_size_formatted": format_size(f.file_size),
        "mime_type": f.mime_type,
        "is_password_protected": share.is_password_protected,
        "view_only": share.view_only,
        "allow_download": share.allow_download,
        "expires_at": to_ist(share.expires_at) if share.expires_at else "Never",
        "views_count": share.view_count,
        "view_count": share.view_count,
        "time_remaining_seconds": int((share.expires_at - datetime.utcnow()).total_seconds()) if share.expires_at else None,
        "download_count": share.download_count,
        "created_at": to_ist_short(share.created_at)
    }

@router.post("/create", response_model=SharedLinkResponse)
def create_shared_link(
    req: CreateSharedLinkRequest,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    if not req.file_id:
        raise HTTPException(status_code=400, detail="file_id is required.")
    return _process_create_share(req.file_id, req, current_user, db)

@router.post("/{file_id}/create", response_model=SharedLinkResponse)
def create_shared_link_by_path(
    file_id: str,
    req: CreateSharedLinkRequest = Body(...),
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    return _process_create_share(file_id, req, current_user, db)

@router.get("/list")
def list_user_shares(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Lists all active shares created by the user."""
    shares = db.query(SharedLink).filter(
        SharedLink.user_id == current_user.id,
        SharedLink.is_active == True
    ).order_by(SharedLink.created_at.desc()).all()

    now = datetime.utcnow()
    res = []
    for s in shares:
        f = s.file
        if not f or f.is_in_recycle_bin:
            continue
        
        is_expired = s.expires_at and s.expires_at < now
        remaining_secs = max(0, int((s.expires_at - now).total_seconds())) if s.expires_at else 86400

        res.append({
            "id": s.id,
            "file_id": f.id,
            "filename": f.filename,
            "file_size": f.file_size,
            "file_size_formatted": format_size(f.file_size),
            "mime_type": f.mime_type,
            "is_password_protected": s.is_password_protected,
            "view_only": s.view_only,
            "allow_download": s.allow_download,
            "expires_at": to_ist(s.expires_at) if s.expires_at else "Never",
            "time_remaining_seconds": remaining_secs,
            "is_expired": is_expired,
            "download_count": s.download_count,
            "view_count": s.view_count,
            "created_at": to_ist_short(s.created_at)
        })
    return res

@router.delete("/{share_id}")
def revoke_share(
    share_id: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Revokes a specific shared link."""
    share = db.query(SharedLink).filter(SharedLink.id == share_id).first()
    if not share:
        raise HTTPException(status_code=404, detail="Shared link not found.")

    if share.user_id != current_user.id and current_user.role != "ADMIN":
        raise HTTPException(status_code=403, detail="Not authorized to revoke this link.")

    share.is_active = False
    db.commit()

    EventService.record_event(
        db, "SHARE_REVOKED", user_id=current_user.id, file_id=share.file_id,
        result="SUCCESS", severity="INFO", metadata={"share_id": share.id}
    )

    AuditService.log(db, "REVOKE_SHARE_LINK", f"Share {share_id}", "SUCCESS", user_id=current_user.id, username=current_user.username)
    return {"message": "Shared link revoked successfully."}

@router.post("/{share_id}/revoke")
def revoke_share_post_alias(
    share_id: str,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    return revoke_share(share_id, current_user, db)

@router.post("/global-revoke")
def global_revoke_all_shares(
    current_admin: User = Depends(get_current_admin),
    db: Session = Depends(get_db)
):
    """Admin feature: Revokes all public shared links globally."""
    active_shares = db.query(SharedLink).filter(SharedLink.is_active == True).all()
    for s in active_shares:
        s.is_active = False
    db.commit()

    AuditService.log(db, "GLOBAL_REVOKE_SHARES", "All Shares", "SUCCESS", "All public links terminated globally", user_id=current_admin.id, username=current_admin.username, role="ADMIN")
    return {"message": f"Successfully revoked {len(active_shares)} active shared link(s)."}

@router.get("/public/{share_id}")
def get_public_share_info(
    share_id: str,
    request: Request,
    db: Session = Depends(get_db)
):
    """Public landing endpoint for shared links."""
    share = db.query(SharedLink).filter(SharedLink.id == share_id, SharedLink.is_active == True).first()
    if not share:
        raise HTTPException(status_code=404, detail="This shared link has expired or been revoked.")

    if share.expires_at and share.expires_at < datetime.utcnow():
        raise HTTPException(status_code=410, detail="This link has expired.")

    f = share.file
    if not f or f.is_in_recycle_bin:
        raise HTTPException(status_code=404, detail="File is no longer available.")

    # 5-second per-IP debounce to prevent double-counting
    client_ip = request.client.host if request.client else "unknown"
    cache_key = f"{share.id}_{client_ip}"
    now_ts = datetime.utcnow().timestamp()
    
    if cache_key not in _VIEW_DEBOUNCE_CACHE or (now_ts - _VIEW_DEBOUNCE_CACHE[cache_key]) > 5.0:
        share.view_count += 1
        db.commit()
        _VIEW_DEBOUNCE_CACHE[cache_key] = now_ts

    return {
        "share_id": share.id,
        "file_id": f.id,
        "filename": f.filename,
        "file_size": f.file_size,
        "file_size_formatted": format_size(f.file_size),
        "mime_type": f.mime_type,
        "extension": f.extension,
        "threat_score": f.threat_score,
        "security_status": f.security_status,
        "is_password_protected": share.is_password_protected,
        "view_only": share.view_only,
        "allow_download": share.allow_download,
        "expires_at": to_ist(share.expires_at) if share.expires_at else "Never",
        "views_count": share.view_count,
        "view_count": share.view_count,
        "created_at": to_ist_short(share.created_at)
    }

@router.post("/public/{share_id}/verify")
def verify_public_share_password(
    share_id: str,
    req: VerifySharePasswordRequest,
    db: Session = Depends(get_db)
):
    share = db.query(SharedLink).filter(SharedLink.id == share_id, SharedLink.is_active == True).first()
    if not share:
        raise HTTPException(status_code=404, detail="This shared link has expired or been revoked.")

    if not share.is_password_protected:
        return {"status": "SUCCESS", "message": "No password required."}

    if not req.password or not share.password_hash or not bcrypt.checkpw(req.password.encode("utf-8"), share.password_hash.encode("utf-8")):
        raise HTTPException(status_code=401, detail="Incorrect password for shared file.")

    return {"status": "SUCCESS", "message": "Password verified successfully."}

@router.get("/public/{share_id}/stream")
def stream_public_shared_file(
    share_id: str,
    db: Session = Depends(get_db)
):
    share = db.query(SharedLink).filter(SharedLink.id == share_id, SharedLink.is_active == True).first()
    if not share or (share.expires_at and share.expires_at < datetime.utcnow()):
        raise HTTPException(status_code=404, detail="Shared link is invalid or expired.")

    f = share.file
    if f: f.storage_path = resolve_storage_path(f.storage_path)
    if not f or not os.path.exists(f.storage_path):
        raise HTTPException(status_code=404, detail="Physical file missing.")

    return FastAPIFileResponse(
        path=f.storage_path,
        media_type=f.mime_type,
        headers={
            "Content-Disposition": f'inline; filename="{f.filename}"',
            "Accept-Ranges": "bytes",
            "Access-Control-Allow-Origin": "*"
        }
    )

@router.get("/public/{share_id}/download")
def download_public_shared_file(
    share_id: str,
    db: Session = Depends(get_db)
):
    share = db.query(SharedLink).filter(SharedLink.id == share_id, SharedLink.is_active == True).first()
    if not share or (share.expires_at and share.expires_at < datetime.utcnow()):
        raise HTTPException(status_code=404, detail="Shared link is invalid or expired.")

    if not share.allow_download or share.view_only:
        raise HTTPException(status_code=403, detail="Download is disabled by the file owner for this link.")

    f = share.file
    if f: f.storage_path = resolve_storage_path(f.storage_path)
    if not f or not os.path.exists(f.storage_path):
        raise HTTPException(status_code=404, detail="Physical file missing.")

    share.download_count += 1
    db.commit()

    return FastAPIFileResponse(
        path=f.storage_path,
        filename=f.filename,
        media_type=f.mime_type
    )

@router.get("/public/{share_id}/view")
def view_public_shared_file_content(
    share_id: str,
    db: Session = Depends(get_db)
):
    import json
    import mammoth
    import docx
    import openpyxl
    import pptx

    share = db.query(SharedLink).filter(SharedLink.id == share_id, SharedLink.is_active == True).first()
    if not share or (share.expires_at and share.expires_at < datetime.utcnow()):
        raise HTTPException(status_code=404, detail="Shared link is invalid or expired.")

    f = share.file
    if f: f.storage_path = resolve_storage_path(f.storage_path)
    if not f or not os.path.exists(f.storage_path):
        raise HTTPException(status_code=404, detail="Physical file missing.")

    ext = f.extension.lower()
    file_size = f.file_size

    with open(f.storage_path, "rb") as fp:
        raw_bytes = fp.read()

    # 1. Plain Text / Code
    if ext in [".txt", ".log", ".md", ".py", ".js", ".ts", ".html", ".css", ".sh", ".bat", ".ps1", ".java", ".c", ".cpp", ".sql", ".ini", ".conf", ".yaml", ".yml"]:
        try:
            text_content = raw_bytes.decode("utf-8", errors="replace")
        except Exception:
            text_content = str(raw_bytes)
        return {
            "format": "TEXT",
            "filename": f.filename,
            "file_size": file_size,
            "mime_type": f.mime_type,
            "content": text_content
        }

    if ext == ".json":
        try:
            parsed = json.loads(raw_bytes.decode("utf-8", errors="replace"))
            formatted = json.dumps(parsed, indent=2)
        except Exception:
            formatted = raw_bytes.decode("utf-8", errors="replace")
        return {
            "format": "JSON",
            "filename": f.filename,
            "file_size": file_size,
            "content": formatted
        }

    if ext == ".csv":
        text_content = raw_bytes.decode("utf-8", errors="replace")
        lines = [line.split(",") for line in text_content.splitlines()]
        return {
            "format": "CSV_TABLE",
            "filename": f.filename,
            "file_size": file_size,
            "rows": lines
        }

    # 2. DOCX Word Document
    if ext in [".docx", ".doc"]:
        html_content = ""
        try:
            mammoth_res = mammoth.convert_to_html(io.BytesIO(raw_bytes))
            html_content = mammoth_res.value
        except Exception:
            pass

        paragraphs = []
        tables = []
        try:
            doc = docx.Document(io.BytesIO(raw_bytes))
            for p in doc.paragraphs:
                if p.text.strip():
                    paragraphs.append({"style": p.style.name if p.style else "Normal", "text": p.text})
            
            for t in doc.tables:
                table_rows = []
                for row in t.rows:
                    table_rows.append([cell.text.strip() for cell in row.cells])
                tables.append(table_rows)
        except Exception:
            pass

        return {
            "format": "DOCX_RENDERED",
            "filename": f.filename,
            "file_size": file_size,
            "html_content": html_content,
            "paragraph_count": len(paragraphs),
            "paragraphs": paragraphs,
            "tables": tables
        }

    # 3. XLSX Excel Sheet
    if ext in [".xlsx", ".xls"]:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
            sheets_data = {}
            for sheetname in wb.sheetnames:
                ws = wb[sheetname]
                rows = []
                for row in ws.iter_rows(values_only=True):
                    if any(c is not None for c in row):
                        rows.append([str(c) if c is not None else "" for c in row])
                    if len(rows) >= 500:
                        break
                sheets_data[sheetname] = rows

            return {
                "format": "XLSX_RENDERED",
                "filename": f.filename,
                "file_size": file_size,
                "sheet_names": wb.sheetnames,
                "sheets": sheets_data
            }
        except Exception as e:
            return {"format": "ERROR", "filename": f.filename, "error": str(e)}

    # 4. PPTX / PPT Presentation (Exact Visual Slide Rendering)
    if ext in [".pptx", ".ppt", ".pps", ".ppsx", ".odp", ".pot", ".potx"]:
        from backend.app.services.presentation_service import PresentationService
        pres_res = PresentationService.render_presentation(raw_bytes, f.file_hash, f.filename)
        pres_res["file_size"] = file_size
        return pres_res

    # 5. PDF, Images, Audio, Video
    if ext in [".pdf", ".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg", ".mp3", ".wav", ".ogg", ".flac", ".mp4", ".webm", ".mov", ".avi"]:
        return {
            "format": "STREAMABLE_MEDIA",
            "filename": f.filename,
            "file_size": file_size,
            "mime_type": f.mime_type,
            "stream_url": f"/api/shares/public/{share.id}/stream"
        }

    return {
        "format": "UNSUPPORTED_BINARY",
        "filename": f.filename,
        "file_size": file_size,
        "file_size_formatted": format_size(file_size),
        "mime_type": f.mime_type,
        "file_hash": f.file_hash,
        "message": "Preview unavailable for this binary format."
    }
